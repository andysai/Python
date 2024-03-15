# 导入库
from docx import Document
from docx.oxml.ns import qn
from setborders import set_cell_border

# 创建word文档
document = Document()

# 创建表格数据
info = [
    ['笔名', '鲁迅'],
    ['原名', '周树人'],
    ['字', '豫才'],
    ['哪里人', '浙江绍兴人'],
    ['第一篇白话小说', '《狂人日记》'],
    ['小说集', '《呐喊》'],
    ['身份1', '文学家'],
    ['身份2', '思想家'],
    ['身份3', '评论家'],
    ['身份4', '作家'],
    ['散文集', '《朝花夕拾》'],
    ['散文诗集', '《野草》']
    ]

# 设置全局字体
document.styles["Normal"].font.name = "宋体"
document.styles["Normal"]._element.rPr.rFonts.set(qn('w:eastAsia'), u'宋体')

# 创建表格
rows = len(info)
cols = len(info[0])
table = document.add_table(rows, cols)

for i in range(rows):
    for j in range(cols):
        table.cell(i, j).text = info[i][j]

        # sz:边界的粗细 val:线性:单线single 虚线dashed 线的颜色:color
        set_cell_border(
            table.cell(i, j),
            top={"sz": 3, "val": "single", "color": "#000000"},
            bottom={"sz": 3, "color": "#000000", "val": "single"},
            left={"sz": 3, "val": "single"},
            right={"sz": 3, "val": "single"},
        )

# 保存文档
document.save("../source_material/01/A/测试(1).docx")
print("执行完成")
