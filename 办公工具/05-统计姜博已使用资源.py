# 导入库
import xlwings as xw
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from setborders import set_cell_border

# 创建实例
app = xw.App(visible=True, add_book=False)
# 打开工作簿
filename, password =r"E:\云达信\设备管理\新机房统计表备份\新机房统计表-20210303.xlsx", "idc123!@#"
workbook = app.books.open(filename, password=password)
# 打开工作表
sheets = workbook.sheets(1)
# 获取总行数列数
sheets_rc = sheets.used_range
nrows = sheets_rc.last_cell.row
# ncols = sheets_rc.last_cell.column

# 数据添加到列表
sheets_row_list = []
for i in range(4, nrows+1):
    first = "$A" + str(i) + ":" + "$AY" + str(i)
    sheets_row_list.append(sheets.range(first).value)

# 获取有墨思达数据的数据
mosida_list = []
# 添加表头
head = ['设备情况', 'CPU', '内存', '硬盘', '存储', '内网IP', '项目']
mosida_list.append(head)

for i in sheets_row_list:
    if i[37] == '深圳市墨思达信息技术有限公司' and i[12] == '虚拟':
        xuniji = ["租用虚拟机", i[19], i[20], i[21], i[22], i[30], i[39]]
        for x in range(7):
            if xuniji[x] == None:
                xuniji[x] = "-"
            else:
                xuniji[x] = xuniji[x]
        mosida_list.append(xuniji)
    elif i[37] == '深圳市墨思达信息技术有限公司' and i[15] == '托管':
        wuliji = ["托管物理机", i[19], i[20], i[21], i[22], i[30], i[39]]
        for n in range(7):
            if wuliji[n] == None:
                wuliji[n] = "-"
            else:
                wuliji[n] = wuliji[n]
        mosida_list.append(wuliji)
# print(mosida_list)
# 关闭工作簿
app.quit()

# 创建word文档
document = Document()

# 设置全局字体
document.styles["Normal"].font.name = "宋体"
document.styles["Normal"]._element.rPr.rFonts.set(qn('w:eastAsia'), u'宋体')
document.styles["Normal"].font.size = Pt(10)
# document.styles["Normal"].font.color.rgb = RGBColor(255, 25, 25)

# 创建表格
rows = len(mosida_list)
cols = len(mosida_list[0])
table = document.add_table(rows, cols)

for i in range(rows):
    for j in range(cols):
        table.cell(i, j).text = mosida_list[i][j]
        # sz:边界的粗细 val:线性:单线single 虚线dashed 线的颜色:color
        set_cell_border(
            table.cell(i, j),
            top={"sz": 3, "val": "single", "color": "#000000"},
            bottom={"sz": 3, "color": "#000000", "val": "single"},
            left={"sz": 3, "val": "single"},
            right={"sz": 3, "val": "single"},
        )

# # 保存文档
document.save(r"E:\云达信\设备管理\新机房统计表备份\姜博设备统计.docx")
print("执行完成")
